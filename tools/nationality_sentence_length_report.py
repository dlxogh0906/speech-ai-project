"""Extend the retained nationality report with country-specific length distributions.

Uses metadata.sqlite without accessing or downloading audio. Counts whitespace
separated eojeol in whole recording labels, not aligned words in audio crops.
Run with the speech-ai Python environment; --out selects an existing report.
"""

import argparse
from collections import Counter, defaultdict
import csv
import datetime
import html
import json
import math
import os
from pathlib import Path
import sqlite3
import zipfile

from nationality_distribution_report import task_group


BINS = [('1-5', 1, 5), ('6-10', 6, 10), ('11-15', 11, 15),
        ('16-20', 16, 20), ('21-30', 21, 30), ('31-50', 31, 50),
        ('51+', 51, None)]
KINDS = {
    'reading_prompt_recording': '낭독 제시문 · 녹음별',
    'reading_prompt_unique': '낭독 제시문 · 중복 문장 제외',
    'spoken_transcript_recording': '실제 발화 전사 · 녹음별',
}
SPLITS = {'train': '학습', 'valid': '검증', 'test': '테스트'}
NOTES = [
    '범위는 기존 보고서와 같은 run_zeroth 마지막 재학습 입력의 41개 모델 국가 라벨이다.',
    '길이는 공백으로 나눈 어절 수다. 형태소 단어 수나 LLM 토큰 수가 아니다.',
    '낭독 제시문은 reading 과제만 집계하며 RW 단어 과제와 질문 제시문을 제외한다.',
    '실제 발화 전사는 단어 과제를 제외한 녹음 중 실제 전사 텍스트가 있는 경우만 집계한다. 질문을 답변 전사로 대체하지 않는다.',
    '전사가 없는 RS·RP 등의 원래 라벨 특성과 원문 미확인 자료는 누락으로 표시한다. 누락된 텍스트를 0어절로 세지 않는다.',
    '구간별 비율은 텍스트가 확인된 집계 대상 수를 분모로 한다. 텍스트 확인률을 함께 제시한다.',
    '녹음별 분포는 같은 문장을 여러 사람이 읽으면 여러 번 센다. 중복 문장 제외 분포는 같은 국가·분할 내 공백 정규화 후 같은 제시문을 한 번 센다.',
    '한 녹음에 여러 문장이 들어갈 수 있다. 문법적 문장 경계로 분리한 분포가 아니라 녹음별 전체 제시문·발화 길이 분포다.',
    '길이는 전체 원문 기준이며 학습 시 음성 크롭 이후 길이는 아니다. 정확한 크롭 구간의 어절 수는 시간 정렬 정보가 필요하다.',
    '차트는 학습 분할이고 CSV·엑셀은 학습·검증·테스트를 모두 포함한다. 41개국 중 해당 분할에 없는 국가는 빈 값으로 남긴다.',
    'P10/P50/P90은 확인된 길이의 누적 순위(nearest-rank)로 계산한다.',
]


def write_json(path, value):
    temporary = path.with_name(path.name + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    os.replace(temporary, path)


def write_csv(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def percentile(counter, q):
    if not counter:
        return None
    target = max(1, math.ceil(sum(counter.values()) * q))
    seen = 0
    for length, count in sorted(counter.items()):
        seen += count
        if seen >= target:
            return length


def new_group():
    return {'reading_eligible': 0, 'transcript_eligible': 0,
            'reading': Counter(), 'transcript': Counter(),
            'prompts': set(), 'transcripts': set()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path,
                        default=Path('/data/speech-ai/out/nationality/data_distribution_20261012'))
    args = parser.parse_args()
    out = args.out
    with (out / 'model_country_split_distribution.csv').open(encoding='utf-8-sig', newline='') as f:
        previous = {(r['label_en'], r['split']): r for r in csv.DictReader(f)}
    countries = {label: row['label_ko'] for (label, _), row in previous.items()}
    groups = defaultdict(new_group)
    connection = sqlite3.connect('file:' + str(out / 'metadata.sqlite') + '?mode=ro', uri=True)
    total = 0
    query = 'SELECT label,split,task,source_task,domain,prompt,transcript FROM report_rows'
    for label, split, task, source_task, domain, prompt, transcript in connection.execute(query):
        total += 1
        g = groups[label, split]
        kind = task_group(source_task or task, domain)
        if kind == 'reading':
            g['reading_eligible'] += 1
            prompt = ' '.join((prompt or '').split())
            if prompt:
                g['reading'][len(prompt.split())] += 1
                g['prompts'].add(prompt)
        if kind != 'word':
            g['transcript_eligible'] += 1
            transcript = ' '.join((transcript or '').split())
            if transcript:
                g['transcript'][len(transcript.split())] += 1
                g['transcripts'].add(transcript)
    connection.close()
    print('[length] aggregated %s model rows' % total, flush=True)

    summaries, bins, exact = [], [], []
    for label in sorted(countries):
        for split in SPLITS:
            g = groups[label, split]
            for kind in KINDS:
                reading = kind.startswith('reading')
                unique = kind.endswith('unique')
                counts = Counter(map(lambda t: len(t.split()), g['prompts'])) if unique else g['reading' if reading else 'transcript']
                eligible = g['reading_eligible' if reading else 'transcript_eligible']
                text_records = sum(g['reading' if reading else 'transcript'].values())
                n = sum(counts.values())
                base = {'label_en': label, 'label_ko': countries[label], 'split': split,
                        'length_type': kind, 'length_type_ko': KINDS[kind]}
                summary = {**base, 'eligible_recordings': eligible,
                           'recordings_with_text': text_records,
                           'recordings_missing_text': eligible - text_records,
                           'text_coverage_pct': round(100 * text_records / eligible, 4) if eligible else None,
                           'distribution_items': n,
                           'unique_texts': len(g['prompts' if reading else 'transcripts']),
                           'mean_eojeol': round(sum(length * c for length, c in counts.items()) / n, 4) if n else None,
                           'p10_eojeol': percentile(counts, .1),
                           'p50_eojeol': percentile(counts, .5),
                           'p90_eojeol': percentile(counts, .9),
                           'min_eojeol': min(counts) if n else None,
                           'max_eojeol': max(counts) if n else None}
                for name, low, high in BINS:
                    count = sum(c for length, c in counts.items() if length >= low and (high is None or length <= high))
                    pct = round(100 * count / n, 4) if n else None
                    summary['bin_' + name + '_count'] = count
                    summary['bin_' + name + '_pct'] = pct
                    bins.append({**base, 'bin_eojeol': name, 'count': count,
                                 'pct_of_known_text': pct, 'distribution_items': n,
                                 'eligible_recordings': eligible,
                                 'recordings_with_text': text_records})
                exact.extend({**base, 'length_eojeol': length, 'count': c,
                              'pct_of_known_text': round(c / n * 100, 6)}
                             for length, c in sorted(counts.items()))
                summaries.append(summary)
    write_csv(out / 'sentence_length_country_summary.csv', summaries)
    write_csv(out / 'sentence_length_country_bins.csv', bins)
    write_csv(out / 'sentence_length_country_exact.csv', exact)

    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.drawing.image import Image

    lookup = {(r['label_en'], r['split'], r['length_type']): r for r in summaries}
    plot_files = []
    for kind, title, filename in [
        ('reading_prompt_recording', 'Reading reference length, weighted by recordings', 'sentence_length_reading_distribution.png'),
        ('spoken_transcript_recording', 'Spoken transcript length, excluding word tasks', 'sentence_length_transcript_distribution.png'),
        ('reading_prompt_unique', 'Reading reference length, counting each text once per country', 'sentence_length_unique_prompt_distribution.png')]:
        labels = sorted(countries)
        data = np.array([[lookup[label, 'train', kind]['bin_' + b[0] + '_pct'] for b in BINS] for label in labels], dtype=float)
        fig, ax = plt.subplots(figsize=(11, 15.4))
        cmap = plt.get_cmap('YlGnBu').copy()
        cmap.set_bad('#e5e7eb')
        chart = ax.imshow(np.ma.masked_invalid(data), vmin=0, vmax=55, cmap=cmap, aspect='auto')
        ax.set_xticks(range(len(BINS)), [b[0] for b in BINS])
        ax.set_yticks(range(len(labels)), [label + '  (n=%s)' % format(lookup[label, 'train', kind]['distribution_items'], ',') for label in labels])
        ax.tick_params(axis='y', labelsize=9)
        ax.set_xlabel('Whitespace-separated Korean eojeol per full recording/reference')
        ax.set_title('Nationality training data: ' + title, fontsize=12, pad=18)
        for i in range(len(labels)):
            for j in range(len(BINS)):
                if not np.isnan(data[i, j]):
                    ax.text(j, i, '%.1f' % data[i, j], ha='center', va='center', fontsize=8,
                            color='white' if data[i, j] > 32 else '#111827')
        fig.colorbar(chart, ax=ax, fraction=.025, pad=.025, label='% of known texts (color capped at 55%)')
        fig.text(.5, .008, 'Train split only. Missing text is excluded, not counted as zero. See coverage in the CSV/Excel.', ha='center', fontsize=8)
        fig.tight_layout(rect=(0, .022, 1, 1))
        fig.savefig(out / filename, dpi=180)
        plt.close(fig)
        plot_files.append(filename)
    print('[length] exported summary, bins, exact counts and three charts', flush=True)

    xlsx = out / '국적분류_데이터분포_20261012.xlsx'
    workbook = load_workbook(xlsx)
    new_names = ['국가별_문장길이', '문장길이_구간별', '문장길이_정확분포', '문장길이_설명', '문장길이_차트']
    for name in new_names:
        if name in workbook:
            del workbook[name]
    sheet = workbook.create_sheet(new_names[0], 1)
    header = ['국가', '영문 라벨', '분할', '길이 기준', '대상 녹음', '텍스트 확인 녹음',
              '텍스트 미확인 녹음', '텍스트 확인률(%)', '분포 집계 수', '고유 원문',
              '평균 어절', 'P10', '중앙값(P50)', 'P90', '최소', '최대']
    header.extend(b[0] + '어절 비율(%)' for b in BINS)
    sheet.append(header)
    for r in summaries:
        keys = ['eligible_recordings', 'recordings_with_text', 'recordings_missing_text',
                'text_coverage_pct', 'distribution_items', 'unique_texts', 'mean_eojeol',
                'p10_eojeol', 'p50_eojeol', 'p90_eojeol', 'min_eojeol', 'max_eojeol']
        sheet.append([r['label_ko'], r['label_en'], SPLITS[r['split']], r['length_type_ko']]
                     + [r[k] for k in keys] + [r['bin_' + b[0] + '_pct'] for b in BINS])
    sheet.conditional_formatting.add('Q2:W%s' % sheet.max_row,
                                    ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF', end_type='num', end_value=55, end_color='2373AB'))
    for title, rows in [(new_names[1], bins), (new_names[2], exact)]:
        sheet = workbook.create_sheet(title)
        sheet.append(list(rows[0]))
        for r in rows:
            sheet.append(list(r.values()))
    sheet = workbook.create_sheet(new_names[3])
    sheet.append(['항목', '집계 기준'])
    for note in NOTES:
        sheet.append(['기준', note])
    sheet = workbook.create_sheet(new_names[4])
    sheet['A1'] = '41개국 학습 분할: 낭독 제시문 / 실제 전사 / 고유 제시문 분포(%)'
    for anchor, filename in zip(['A3', 'A77', 'A151'], plot_files):
        img = Image(out / filename)
        img.width, img.height = 792, 1109
        sheet.add_image(img, anchor)
    for name in new_names[:-1]:
        sheet = workbook[name]
        sheet.freeze_panes = 'E2' if name == new_names[0] else 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='23405F')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26 if column[0].column < 5 else 21
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '0.00'
    workbook[new_names[3]].column_dimensions['B'].width = 110
    temporary = xlsx.with_name(xlsx.stem + '.tmp.xlsx')
    workbook.save(temporary)
    os.replace(temporary, xlsx)

    report = out / '국적분류_데이터분포_20261012.html'
    body = report.read_text()
    begin, end = '<!-- sentence-length:start -->', '<!-- sentence-length:end -->'
    if begin in body:
        first, tail = body.split(begin, 1)
        _, last = tail.split(end, 1)
        body = first + last
    sections = [begin, '<h2>국가별 녹음·문장 길이 분포</h2>',
                '<p>학습 분할의 확인된 원문을 1–5 / 6–10 / 11–15 / 16–20 / 21–30 / 31–50 / 51+ 어절로 나눴습니다. 단어 과제는 제외했습니다. 누락 원문은 분모에서 제외하고, 엑셀에 확인률을 별도로 표시했습니다.</p>']
    for filename, caption in zip(plot_files, ['낭독 제시문 · 녹음별', '실제 발화 전사 · 녹음별', '낭독 제시문 · 중복 문장 제외']):
        sections.append('<details%s><summary>%s</summary><img src="%s" alt="%s" style="width:min(100%%,1000px);height:auto"></details>' % (' open' if filename == plot_files[0] else '', caption, filename, caption))
    sections.append('<h3>우즈베키스탄·필리핀 낭독 제시문 비교</h3><table><tr><th>라벨</th><th>확인 녹음</th><th>평균</th><th>중앙값</th><th>P90</th>' + ''.join('<th>' + b[0] + '어절(%)</th>' for b in BINS) + '</tr>')
    for label in ['Uzbekistan', 'Philippines']:
        r = lookup[label, 'train', 'reading_prompt_recording']
        values = [countries[label], r['distribution_items'], r['mean_eojeol'], r['p50_eojeol'], r['p90_eojeol']] + [round(r['bin_' + b[0] + '_pct'], 2) for b in BINS]
        sections.append('<tr>' + ''.join('<td>' + html.escape(str(v)) + '</td>' for v in values) + '</tr>')
    sections.append('</table><ul>' + ''.join('<li>' + html.escape(n) + '</li>' for n in NOTES) + '</ul>' + end)
    report.write_text(body.replace('</html>', '\n'.join(sections) + '\n</html>'))

    summary_path = out / 'sentence_length_summary.json'
    summary = {'generated_at': datetime.datetime.now().astimezone().isoformat(),
               'source': str(out / 'metadata.sqlite'), 'model_manifest_rows': total,
               'countries': len(countries), 'bins_eojeol': [b[0] for b in BINS],
               'length_types': KINDS, 'notes': NOTES,
               'priority_countries_train': {label: {kind: lookup[label, 'train', kind] for kind in KINDS} for label in ['Uzbekistan', 'Philippines', 'Russia', 'Korea']},
               'country_split_type_summaries': len(summaries),
               'plots': plot_files}
    write_json(summary_path, summary)
    main_summary = json.loads((out / 'summary.json').read_text())
    main_summary['sentence_length_report'] = {'summary': summary_path.name, 'bins_eojeol': summary['bins_eojeol'], 'length_types': KINDS, 'countries': len(countries), 'plots': plot_files}
    write_json(out / 'summary.json', main_summary)
    draft = out / '대표님_전달문.txt'
    draft_text = draft.read_text()
    paragraph = '\n국가별 문장·발화 길이 분포도 추가했습니다. 어절 수 7개 구간별 건수·비율과 평균·중앙값·P90, 텍스트 확인률을 제공합니다. 낭독 제시문, 실제 발화 전사, 중복 제시문 제외 분포를 분리했으며 학습·검증·테스트별로 조회할 수 있습니다. 학습 41개국 분포 그래프 3개도 함께 첨부했습니다.\n'
    if '국가별 문장·발화 길이 분포도 추가했습니다.' not in draft_text:
        draft.write_text(draft_text + paragraph)
    bundle = out / '국적분류_데이터분포_전달용_20261012.zip'
    temporary = bundle.with_name(bundle.name + '.tmp')
    with zipfile.ZipFile(temporary, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=3) as archive:
        for p in sorted(out.iterdir()):
            if p.suffix in ('.csv', '.xlsx', '.html', '.png') or p.name.endswith('.csv.gz') or p.name in ('summary.json', 'provenance.json', 'sentence_length_summary.json', '대표님_전달문.txt'):
                archive.write(p, p.name)
    os.replace(temporary, bundle)
    print('[complete] updated Excel, HTML, handoff ZIP and text draft', flush=True)
    print(json.dumps(summary['priority_countries_train'], ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
