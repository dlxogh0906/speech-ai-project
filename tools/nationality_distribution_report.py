"""Reconstruct nationality data distribution and transcripts from retained metadata.

The final run_zeroth manifests are checked against its actual loader CSVs.
Historical inventory counts come from retained pre-deletion corpus indexes.
Audio availability is measured separately; no audio is downloaded or inferred.
"""
import argparse
import collections
import csv
import datetime
import gzip
import hashlib
import html
import json
import multiprocessing
import os
from pathlib import Path
import re
import sqlite3
import time
import zipfile

REPO = Path(__file__).resolve().parents[1]
NAT = Path('/data/speech-ai/out/nationality')
BASE = NAT / 'webmic_kd'
CORPUS = Path('/data/speech-ai/corpus')
FREE = {'EIG', 'PDT', 'SPT', 'SSO', 'KOR'}


def write_json(path, value):
    path = Path(path); temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    os.replace(temporary, path)


def normalize(value):
    return ' '.join(str(value).split()) if isinstance(value, str) else ''


def task_group(task, domain=''):
    if task == 'RW': return 'word'
    if task in ('RS', 'LAR', 'RP') or task.startswith('R') and len(task) == 2: return 'reading'
    if task == 'ATQ' or task.startswith('Q') and len(task) == 2: return 'question_answer'
    if task in FREE or domain in ('kspon', 'kor001', 'kor_aug_train'): return 'free'
    return 'other'


def source_key(row):
    domain = row['domain']; stem = Path(row['audio_path']).stem
    if domain.startswith('web'): dataset = domain[3:]
    elif domain == 'p505': dataset = '131'
    elif domain == 'kspon' or stem.startswith('KsponSpeech_'): dataset = 'kspon'
    elif domain.startswith('zeroth'): dataset = 'zeroth'
    else: dataset = '001'
    return dataset, re.sub(r'_aug\d+$', '', stem)


def metadata_job(job):
    dataset, stem, path = job
    try:
        try: data = json.loads(Path(path).read_text(encoding='utf-8-sig'))
        except UnicodeDecodeError: data = json.loads(Path(path).read_text(encoding='cp949'))
        if dataset in ('040', '041', '042', '043'):
            rec = data.get('RecordingMetadata', {})
            transcript, prompt = normalize(rec.get('orthographic')), normalize(rec.get('prompt'))
            task = re.sub(r'\d+', '', stem.split('-')[5])
            meta = data.get('SpeakerMetadata', {})
            country, language = str(meta.get('nationality', '')), str(meta.get('language', ''))
            method = 'RecordingMetadata.orthographic'
        elif dataset == '131':
            rec = data.get('transcription', {})
            task = stem[4:6]
            transcript = normalize(rec.get('ReadingLabelText') if task.startswith('R') else rec.get('AnswerLabelText'))
            prompt = normalize(rec.get('Reading') if task.startswith('R') else rec.get('Question'))
            country = str(data.get('residence_info', {}).get('country', ''))
            language = str(data.get('skill_info', {}).get('motherTongue', ''))
            method = 'transcription.' + ('ReadingLabelText' if task.startswith('R') else 'AnswerLabelText')
        else:
            return dataset, stem, '', '', '', '', '', 'unsupported_native_json', ''
        return dataset, stem, transcript, prompt, task, country, language, method, ''
    except Exception as exc:
        return dataset, stem, '', '', '', '', '', 'label_parse_error', type(exc).__name__


def write_csv(path, rows):
    if not rows: return
    with Path(path).open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', default=str(NAT / 'data_distribution_20261012'))
    parser.add_argument('--workers', type=int, default=6)
    args = parser.parse_args(); out = Path(args.out); out.mkdir(parents=True, exist_ok=True)
    started = time.time()
    def progress(stage, **extra):
        write_json(out / 'status.json', {'status': 'running', 'stage': stage, 'updated_at': datetime.datetime.now().astimezone().isoformat(), **extra})
        print('[distribution] ' + json.dumps({'stage': stage, **extra}, ensure_ascii=False), flush=True)
    connection = sqlite3.connect(str(out / 'metadata.sqlite'))
    connection.execute('PRAGMA journal_mode=WAL')
    connection.execute('CREATE TABLE IF NOT EXISTS rows(id INTEGER PRIMARY KEY,dataset TEXT,stem TEXT,split TEXT,label TEXT,speaker TEXT,domain TEXT,task TEXT,duration REAL,audio_path TEXT,audio_exists INTEGER,augmented INTEGER)')
    connection.execute('CREATE TABLE IF NOT EXISTS annotations(dataset TEXT,stem TEXT,transcript TEXT,prompt TEXT,task TEXT,metadata_country TEXT,language TEXT,method TEXT,error TEXT,PRIMARY KEY(dataset,stem))')
    connection.execute('CREATE TABLE IF NOT EXISTS stages(name TEXT PRIMARY KEY)')
    mapping = json.loads((REPO / 'tasks/nationality-classification/configs/country41_ko2en.json').read_text())
    ko = {v: k for k, v in mapping.items()}; ko['Korea'] = '한국'
    provenance = {'generated_at': datetime.datetime.now().astimezone().isoformat(), 'model': 'run_zeroth final training stage', 'input_files': {}}

    def digest_rows(path, label_key):
        hashed = hashlib.sha256(); count = 0
        with Path(path).open(newline='') as f:
            for row in csv.DictReader(f):
                hashed.update(json.dumps([row['audio_path'], row[label_key]], ensure_ascii=False, separators=(',', ':')).encode()); hashed.update(b'\n')
                count += 1
        return count, hashed.hexdigest()

    progress('verify_training_loader')
    for split in ['train', 'valid', 'test']:
        path = BASE / 'lab131full_zeroth' / (split + '.csv')
        count, digest = digest_rows(path, 'label_en')
        provenance['input_files'][str(path)] = {'rows': count, 'ordered_audio_path_label_sha256': digest, 'bytes': path.stat().st_size}
        if split != 'test':
            loader = BASE / 'run_zeroth' / ('_' + split + '_loader.csv')
            loader_count, loader_digest = digest_rows(loader, 'label_name')
            if (count, digest) != (loader_count, loader_digest): raise RuntimeError('Model loader differs from ' + split + ' manifest')
            provenance['input_files'][str(loader)] = {'rows': loader_count, 'ordered_audio_path_label_sha256': loader_digest, 'matches_manifest': True}
    if not connection.execute("SELECT 1 FROM stages WHERE name='manifests'").fetchone():
        progress('import_model_manifests')
        identifier, pending = 0, []
        for split in ['train', 'valid', 'test']:
            with (BASE / 'lab131full_zeroth' / (split + '.csv')).open(newline='') as f:
                for row in csv.DictReader(f):
                    identifier += 1; dataset, stem = source_key(row)
                    pending.append((identifier, dataset, stem, split, row['label_en'], row['speaker'], row['domain'], row['task'], float(row['duration_sec']), row['audio_path'], int(Path(row['audio_path']).is_file()), int('aug' in row['domain'])))
                    if len(pending) >= 5000:
                        connection.executemany('INSERT OR REPLACE INTO rows VALUES (?,?,?,?,?,?,?,?,?,?,?,?)', pending); pending.clear()
            if pending: connection.executemany('INSERT OR REPLACE INTO rows VALUES (?,?,?,?,?,?,?,?,?,?,?,?)', pending); pending.clear()
            connection.commit(); progress('import_model_manifests', split=split, cumulative_rows=identifier)
        connection.execute('CREATE INDEX IF NOT EXISTS row_source ON rows(dataset,stem)')
        connection.execute('CREATE INDEX IF NOT EXISTS row_class_split ON rows(label,split)')
        connection.execute("INSERT INTO stages VALUES ('manifests')"); connection.commit()
    wanted = collections.defaultdict(set)
    for dataset, stem in connection.execute('SELECT DISTINCT dataset,stem FROM rows'): wanted[dataset].add(stem)
    existing = {(d, s) for d, s in connection.execute('SELECT dataset,stem FROM annotations')}

    progress('read_retained_labels', selected_source_records=sum(map(len, wanted.values())), already_read=len(existing))
    def jobs():
        for dataset in ['040', '041', '042', '043', '131']:
            base = next((CORPUS / 'aihub').glob(dataset + '.*'))
            for directory, subdirs, filenames in os.walk(base):
                if '라벨링' not in directory:
                    subdirs[:] = [s for s in subdirs if not ('원천데이터' in s)]
                    continue
                for filename in filenames:
                    if filename.endswith('.json'):
                        stem = filename[:-5]
                        if stem in wanted[dataset] and (dataset, stem) not in existing:
                            yield dataset, stem, str(Path(directory) / filename)
    pending, count = [], 0
    with multiprocessing.Pool(args.workers) as pool:
        for result in pool.imap_unordered(metadata_job, jobs(), chunksize=128):
            pending.append(result); count += 1
            if len(pending) >= 2000:
                connection.executemany('INSERT OR REPLACE INTO annotations VALUES (?,?,?,?,?,?,?,?,?)', pending); pending.clear(); connection.commit()
            if count % 20000 == 0: progress('read_retained_labels', labels_read=count, elapsed_seconds=round(time.time() - started))
    if pending: connection.executemany('INSERT OR REPLACE INTO annotations VALUES (?,?,?,?,?,?,?,?,?)', pending); connection.commit()

    progress('native_transcripts')
    native = []
    for path in (CORPUS / 'zeroth').rglob('*.trans.txt'):
        for line in path.read_text().splitlines():
            parts = line.split(maxsplit=1)
            if len(parts) == 2 and parts[0] in wanted['zeroth']:
                text = normalize(parts[1]); native.append(('zeroth', parts[0], text, text, 'RS', 'Korea', 'Korean', 'Zeroth.trans.txt', ''))
    scripts = CORPUS / 'aihub/10.한국어음성/KsponSpeech_scripts'
    for path in scripts.glob('*.trn'):
        with path.open(encoding='utf-8-sig') as f:
            for line in f:
                parts = line.split(' :: ', 1)
                if len(parts) != 2: continue
                stem = Path(parts[0]).stem
                if stem not in wanted['kspon']: continue
                text = re.sub(r'\(([^()]*)\)/\(([^()]*)\)', lambda m: m.group(2), parts[1])
                text = re.sub(r'(^|\s)[bnlo]/(?=\s|$)', ' ', text)
                text = re.sub(r'([가-힣A-Za-z0-9]+)/', r'\1', text).replace('+', ' ')
                native.append(('kspon', stem, normalize(text), '', 'KOR', 'Korea', 'Korean', 'KsponSpeech.phonetic_transcript', ''))
    connection.executemany('INSERT OR REPLACE INTO annotations VALUES (?,?,?,?,?,?,?,?,?)', native); connection.commit()

    progress('historical_inventory')
    pool_stats, source_meta = {}, {}
    for name in ['index_040_043.csv', 'index_131_labeled.csv']:
        path = BASE / name; count = 0
        with path.open(newline='') as f:
            for row in csv.DictReader(f):
                count += 1; dataset = row.get('dataset', '131'); label = row['label_en'] or '__outside_model__'
                metadata_country = row.get('nationality', '') if dataset != '131' else row['country']
                key = dataset, label, metadata_country
                group = pool_stats.setdefault(key, {'rows': 0, 'speakers': set(), 'tasks': collections.Counter(), 'seconds': 0.})
                task = row.get('task') or row['stem'][4:6]
                group['rows'] += 1; group['speakers'].add(row['speaker']); group['tasks'][task_group(task)] += 1
                group['seconds'] += float(row.get('duration_sec') or 0)
        source_meta[name] = {'rows': count, 'path': str(path), 'bytes': path.stat().st_size, 'scope': 'Retained historical pre-deletion inventory, not current audio availability'}
    pool_rows = []
    for (dataset, label, metadata_country), g in sorted(pool_stats.items()):
        pool_rows.append({'dataset': dataset, 'model_label': label, 'label_ko': ko.get(label, ''),
                          'metadata_country': metadata_country, 'metadata_country_kind': 'nationality' if dataset != '131' else 'residence_country',
                          'recordings_historical': g['rows'], 'speaker_keys': len(g['speakers']), 'hours_annotated': round(g['seconds'] / 3600, 3),
                          **{k + '_records': g['tasks'][k] for k in ['word', 'reading', 'question_answer', 'free', 'other']}})
    write_csv(out / 'historical_corpus_distribution.csv', pool_rows)

    progress('aggregate_transcripts')
    connection.execute('DROP VIEW IF EXISTS report_rows')
    connection.execute('CREATE VIEW report_rows AS SELECT r.*,a.transcript,a.prompt,a.task AS source_task,a.metadata_country,a.language,a.method,a.error FROM rows r LEFT JOIN annotations a ON r.dataset=a.dataset AND r.stem=a.stem')
    groups = {}; examples = []; example_counts = collections.Counter(); prompt_counts = collections.Counter(); utterance_count = 0
    sentence_fields = ['record_id', 'split', 'model_label', 'label_ko', 'dataset', 'domain', 'speaker_anonymous_key', 'speaker_key_kind', 'task_manifest', 'task_source', 'task_group', 'duration_sec', 'augmented', 'audio_currently_present', 'transcript', 'transcript_eojeol_count', 'prompt', 'prompt_eojeol_count', 'text_source', 'metadata_country', 'metadata_country_kind', 'mother_tongue']
    with gzip.open(out / 'utterances_with_word_counts.csv.gz', 'wt', encoding='utf-8-sig', newline='', compresslevel=3) as sink:
        writer = csv.DictWriter(sink, fieldnames=sentence_fields); writer.writeheader()
        query = 'SELECT label,split,speaker,dataset,stem,domain,task,duration,audio_exists,augmented,transcript,prompt,source_task,method,metadata_country,language FROM report_rows ORDER BY id'
        for label, split, speaker, dataset, stem, domain, task, duration, audio_exists, augmented, transcript, prompt, actual_task, method, metadata_country, language in connection.execute(query):
            transcript, prompt = transcript or '', prompt or ''; actual_task = actual_task or task
            typ = task_group(actual_task, domain); key = label, split
            g = groups.setdefault(key, {'rows': 0, 'speakers': set(), 'verified_keys': set(), 'group_keys': set(), 'base_records': set(), 'transcripts': set(), 'prompts': set(), 'reading_prompts': set(), 'seconds': 0., 'augmented': 0, 'audio_present': 0, 'word_counts': collections.Counter(), 'prompt_word_counts': collections.Counter(), 'reading_prompt_word_counts': collections.Counter(), 'transcript_rows': 0, 'prompt_rows': 0, 'types': collections.Counter(), 'domains': collections.Counter()})
            g['rows'] += 1; g['speakers'].add(speaker); g['base_records'].add((dataset, stem)); g['seconds'] += duration
            kind = 'folder_group' if speaker.startswith('KSPON:') else 'provider_speaker_key'
            g['group_keys' if kind == 'folder_group' else 'verified_keys'].add(speaker)
            g['augmented'] += augmented; g['audio_present'] += audio_exists; g['types'][typ] += 1; g['domains'][domain] += 1
            wc, pc = len(transcript.split()), len(prompt.split())
            if transcript: g['transcript_rows'] += 1; g['word_counts'][wc] += 1; g['transcripts'].add(transcript)
            if prompt:
                g['prompt_rows'] += 1; g['prompt_word_counts'][pc] += 1; g['prompts'].add(prompt)
                if typ == 'reading': g['reading_prompt_word_counts'][pc] += 1; g['reading_prompts'].add(prompt)
                if split == 'train': prompt_counts[(label, typ, prompt, pc)] += 1
            record = {'record_id': hashlib.sha256((str(utterance_count + 1) + ':' + dataset + ':' + stem + ':' + domain).encode()).hexdigest()[:24],
                      'split': split, 'model_label': label, 'label_ko': ko.get(label, label), 'dataset': dataset, 'domain': domain,
                      'speaker_anonymous_key': hashlib.sha256(('nationality-report:' + speaker).encode()).hexdigest()[:24],
                      'speaker_key_kind': kind, 'task_manifest': task, 'task_source': actual_task, 'task_group': typ,
                      'duration_sec': duration, 'augmented': augmented, 'audio_currently_present': audio_exists,
                      'transcript': transcript, 'transcript_eojeol_count': wc if transcript else '',
                      'prompt': prompt, 'prompt_eojeol_count': pc if prompt else '', 'text_source': method or 'unavailable',
                      'metadata_country': metadata_country or '', 'metadata_country_kind': 'residence_country' if dataset == '131' else 'nationality_or_native_corpus',
                      'mother_tongue': language or ''}
            writer.writerow(record); utterance_count += 1
            if split == 'train' and transcript and example_counts[label] < 12:
                examples.append(record); example_counts[label] += 1
    def stats(counter):
        n = sum(counter.values())
        if not n: return {'mean': None, 'p50': None, 'min': None, 'max': None, 'sum': 0}
        cumulative, median = 0, None
        for length, count in sorted(counter.items()):
            cumulative += count
            if median is None and cumulative >= (n + 1) / 2: median = length
        return {'mean': round(sum(k*v for k,v in counter.items()) / n, 3), 'p50': median, 'min': min(counter), 'max': max(counter), 'sum': sum(k*v for k,v in counter.items())}
    class_rows = []; speaker_sets = collections.defaultdict(set)
    for (label, split), g in sorted(groups.items()):
        word_stats, prompt_stats, reading_stats = stats(g['word_counts']), stats(g['prompt_word_counts']), stats(g['reading_prompt_word_counts'])
        speaker_sets[split].update(g['speakers'])
        class_rows.append({'label_en': label, 'label_ko': ko.get(label, label), 'split': split, 'speaker_keys': len(g['speakers']),
                           'provider_speaker_keys': len(g['verified_keys']), 'folder_group_keys_not_people': len(g['group_keys']),
                           'data_rows': g['rows'], 'unique_source_recordings': len(g['base_records']), 'augmentation_rows': g['augmented'],
                           'hours_annotated': round(g['seconds'] / 3600, 3), 'audio_present_now': g['audio_present'],
                           'audio_missing_now': g['rows'] - g['audio_present'],
                           **{k + '_records': g['types'][k] for k in ['word', 'reading', 'question_answer', 'free', 'other']},
                           'sentence_type_records': g['types']['reading'] + g['types']['question_answer'],
                           'transcript_available_records': g['transcript_rows'], 'transcript_coverage_pct': round(100*g['transcript_rows']/g['rows'], 2),
                           'unique_transcripts': len(g['transcripts']), 'transcript_eojeol_total': word_stats['sum'],
                           'transcript_eojeol_mean': word_stats['mean'], 'transcript_eojeol_p50': word_stats['p50'],
                           'transcript_eojeol_min': word_stats['min'], 'transcript_eojeol_max': word_stats['max'],
                           'prompt_available_records': g['prompt_rows'], 'unique_prompts': len(g['prompts']), 'prompt_eojeol_mean': prompt_stats['mean'],
                           'reading_prompt_available_records': sum(g['reading_prompt_word_counts'].values()),
                           'unique_reading_prompts': len(g['reading_prompts']), 'reading_prompt_eojeol_mean': reading_stats['mean'],
                           'reading_prompt_eojeol_p50': reading_stats['p50'], 'reading_prompt_eojeol_min': reading_stats['min'],
                           'reading_prompt_eojeol_max': reading_stats['max']})
    write_csv(out / 'model_country_split_distribution.csv', class_rows)
    write_csv(out / 'sentence_examples.csv', examples)
    prompts = [{'label_en': l, 'label_ko': ko.get(l,l), 'task_group': t, 'prompt': p, 'prompt_eojeol_count': n, 'training_records_with_this_prompt': count} for (l,t,p,n),count in sorted(prompt_counts.items())]
    write_csv(out / 'training_prompts_word_counts.csv', prompts)
    domains = []
    for label, split, domain, count, people, seconds, available in connection.execute('SELECT label,split,domain,COUNT(*),COUNT(DISTINCT speaker),SUM(duration),SUM(audio_exists) FROM rows GROUP BY label,split,domain'):
        domains.append({'label_en': label, 'label_ko': ko.get(label,label), 'split': split, 'domain': domain, 'rows': count, 'speaker_keys': people, 'hours_annotated': round(seconds/3600,3), 'audio_present_now': available})
    write_csv(out / 'model_country_domain_distribution.csv', domains)
    mother_tongues = []
    for label, split, dataset, language, count, people in connection.execute('SELECT label,split,dataset,COALESCE(language,"unknown"),COUNT(*),COUNT(DISTINCT speaker) FROM report_rows GROUP BY label,split,dataset,language'):
        mother_tongues.append({'label_en':label,'label_ko':ko.get(label,label),'split':split,'dataset':dataset,
                              'mother_tongue_raw':language,'data_rows':count,'speaker_keys':people})
    write_csv(out / 'model_country_mother_tongue_distribution.csv', mother_tongues)
    training = [r for r in class_rows if r['split'] == 'train']; by_country = {r['label_en']: r for r in training}
    counts = [r['data_rows'] for r in training]; people = [r['speaker_keys'] for r in training]
    notes = [
        '기준 모델은 NPU용 run_zeroth의 마지막 재학습 단계다. 실제 _train_loader.csv와 입력 순서·경로·라벨 해시가 일치한다.',
        '이전 교사·초기 모델의 학습 이력을 합산한 누적 고유 데이터 수는 이 표의 범위에 포함되지 않는다.',
        '원본 수집량은 삭제 전 인덱스의 역사적 녹음 수다. 현재 남아 있는 음성 수는 별도 열에 표시한다.',
        '국가별 모델 라벨과 실제 국적은 동일한 개념이 아니다. 040~043은 nationality, 131은 모국어와 거주국을 기존 규칙으로 국가 라벨에 매핑한다.',
        '131의 metadata_country는 residence_info.country(거주국)이며, 실제 국적을 나타내는 필드로 해석하지 않는다.',
        '화자 수는 데이터셋 namespace를 포함한 학습 매니페스트의 고유 화자키 수다. KSPON 키는 폴더 그룹으로 실제 사람 수가 확인되지 않는다.',
        '문장형 녹음 수는 낭독과 질문 응답 파일 수다. 한 녹음에 여러 문장이 들어갈 수 있다. 단어 과제와 자유발화는 별도 열이다.',
        '어절 수는 공백으로 구분한 원문 단위 수다. 형태소 단어 수·LLM 토큰 수와 구분한다. 원문 누락은 빈칸으로 표시하고 평균에서 제외한다.',
        'prompt는 제시문·질문, transcript는 실제 발화 전사다. 질문의 단어 수를 답변 발화의 단어 수로 대체하지 않는다.',
        'RS·RW·RP 발음 과제는 원래 실제 발화 전사 없이 제시문만 제공하는 경우가 많다. 낭독 제시문의 어절 수와 실제 전사의 어절 수를 각각 제공한다.',
        '문장 텍스트와 어절 수는 전체 원본 라벨 기준이다. 학습 음성의 시간 제한·크롭 이후 단어 수는 시간 정렬이 필요하다.',
        '한국어 40,000행은 기존 녹음 증강본이다. 증강 행 수와 고유 원본 녹음 수를 함께 표시한다.',
        '001 자유대화 라벨은 현재 경로에서 찾지 못한 경우 원문 미확인으로 표시한다. 이 집계는 없는 원문을 만들어 채우지 않는다.',
        '배포 주 모델의 마지막 단계 분포를 집계했으며 별도 게이트 모델, 이전 A100 모델, YouTube·LoRA 추가 실험을 이 수에 합산하지 않는다.'
    ]
    deletion = json.loads(Path('/data/speech-ai/out/cleanup/nationality_20261011/deletion_report.json').read_text())
    summary = {'model_scope': provenance['model'], 'totals': {s: {'data_rows': sum(r['data_rows'] for r in class_rows if r['split']==s), 'speaker_keys': len(speaker_sets[s]), 'classes_with_data': sum(r['split']==s for r in class_rows)} for s in ['train','valid','test']},
               'speaker_key_overlap': {a+'_'+b:len(speaker_sets[a]&speaker_sets[b]) for a,b in [('train','valid'),('train','test'),('valid','test')]},
               'historical_foreign_inventory_rows': sum(r['recordings_historical'] for r in pool_rows),
               'class_imbalance_training_max_min_rows': round(max(counts)/min(counts),2),
               'class_imbalance_training_max_min_speaker_keys': round(max(people)/min(people),2),
               'top5_training_row_share_pct': round(sum(sorted(counts,reverse=True)[:5])/sum(counts)*100,2),
               'small_train_classes_under_20_keys': [r for r in training if r['speaker_keys']<20],
               'priority_countries': {label: by_country[label] for label in ['Uzbekistan','Philippines','Russia','Korea']},
               'transcript_coverage': {s: {'available':sum(r['transcript_available_records'] for r in class_rows if r['split']==s),'total':sum(r['data_rows'] for r in class_rows if r['split']==s)} for s in ['train','valid','test']},
               'prompt_coverage': {s: {'available':sum(r['prompt_available_records'] for r in class_rows if r['split']==s),'total':sum(r['data_rows'] for r in class_rows if r['split']==s)} for s in ['train','valid','test']},
               'notes': notes, 'deletion': deletion, 'source_inventory':source_meta}
    write_json(out / 'summary.json', summary); write_json(out / 'provenance.json', provenance)

    progress('render_shareable_files')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    workbook = Workbook(); sheet = workbook.active; sheet.title = '설명'
    sheet.append(['항목','설명']); sheet.append(['모델·집계 범위',provenance['model']]); sheet.append(['학습 행 수',summary['totals']['train']['data_rows']])
    for note in notes: sheet.append(['집계 기준',note])
    for title, values in [('국가별_학습',training),('국가별_분할',class_rows),('원본수집량',pool_rows),('국가별_도메인',domains),('국가별_모국어',mother_tongues),('문장예시',examples),('제시문_어절수',prompts)]:
        if not values: continue
        s=workbook.create_sheet(title);s.append(list(values[0]));
        for row in values: s.append([row.get(k) for k in values[0]])
    for s in workbook:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for cell in s[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23405F')
        for column in s.columns:
            name=column[0].value; s.column_dimensions[column[0].column_letter].width=65 if name in ('prompt','transcript','설명') else 23
        for row in s.iter_rows():
            for cell in row:
                if cell.data_type=='f': cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
    workbook.save(out/'국적분류_데이터분포_20261012.xlsx')
    body=['<!doctype html><html lang="ko"><meta charset="utf-8"><title>국적 분류 데이터 분포</title>',
          '<style>body{font:16px system-ui;color:#172435;max-width:1450px;margin:32px auto;padding:0 20px}table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:9px;border:1px solid #d6dfe8;text-align:right}th{background:#e9eff6}td:first-child,th:first-child{text-align:left}li,p{line-height:1.65}.lead{background:#edf5ff;padding:18px}h2{margin-top:32px}</style>',
          '<h1>국적 분류 데이터 분포 · 문장 및 어절 수</h1>',
          '<p class="lead">run_zeroth 마지막 재학습 단계: <strong>%s행</strong> · 41개 라벨. 실제 학습 로더와 입력 목록의 해시 일치 확인.</p>'%format(summary['totals']['train']['data_rows'],','),
          '<p>삭제 전 외국인 데이터 인덱스 %s행이 남아 있어 분포를 재구성했습니다. 국적 전용 음성과 특징 캐시 삭제로 실제 약 %.1fGiB를 확보했으며 라벨·문장·모델·학습 목록은 보존됐습니다.</p>'%(format(summary['historical_foreign_inventory_rows'],','),deletion['disk_free_change']/2**30),
          '<h2>국가별 마지막 단계 학습 데이터</h2><table><tr><th>모델 라벨</th><th>화자키</th><th>학습 행</th><th>단어 과제</th><th>낭독</th><th>질문 응답</th><th>자유발화</th><th>고유 낭독 제시문</th><th>낭독 평균 어절</th><th>전사 평균 어절</th></tr>']
    for r in sorted(training,key=lambda r:-r['data_rows']):
        values=[r['label_ko']+' ('+r['label_en']+')',r['speaker_keys'],r['data_rows'],r['word_records'],r['reading_records'],r['question_answer_records'],r['free_records'],r['unique_reading_prompts'],r['reading_prompt_eojeol_mean'],r['transcript_eojeol_mean']]
        body.append('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in values)+'</tr>')
    body.append('</table><h2>집계 기준과 해석 범위</h2><ul>'+''.join('<li>'+html.escape(n)+'</li>' for n in notes)+'</ul></html>')
    (out/'국적분류_데이터분포_20261012.html').write_text('\n'.join(body),encoding='utf-8')
    bundle=out/'국적분류_데이터분포_전달용_20261012.zip'
    with zipfile.ZipFile(bundle,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=3) as archive:
        for path in out.iterdir():
            if path.suffix in ('.csv','.xlsx','.html') or path.name.endswith('.csv.gz') or path.name in ('summary.json','provenance.json'):
                archive.write(path,path.name)
    connection.close()
    write_json(out/'status.json',{'status':'complete','finished_at':datetime.datetime.now().astimezone().isoformat(),'elapsed_seconds':round(time.time()-started,1),'utterance_rows':utterance_count,'bundle':str(bundle)})
    print('[complete] '+json.dumps(summary['totals']),flush=True)


if __name__=='__main__': main()
